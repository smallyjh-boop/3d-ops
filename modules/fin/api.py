"""电商财务（v2.21；v2.30 起是「财务」模块）的接口。只调用 logic 的公开函数。网页接口只有老板能用（auth.OWNER_ONLY）；
对外接口 /api/v1/finance/…（密钥权限 finance，只读）给以后的财务软件 / 会计用。"""
import io

from flask import Blueprint

from web.common import *  # noqa: F401,F403
from modules.fin import logic as fin

bp = Blueprint("m_fin", __name__)
PREFIXES = ["/api/fin/", "/api/expenses"] + [f"/api/v1/finance/{x}" for x in ("pl", "breakdown", "cashflow", "expenses", "tax")]
fin.init()          # 自己的数据表


def _method_zh(m):
    """v2.41：导出里付款方式写中文（以前是 transfer / card 代码）。"""
    from modules.fin.logic import book
    return book.METHODS.get(m or "", m or "")


def _period():
    a = request.args
    try:
        if a.get("start"):
            return fin.period(start=a.get("start"), end=a.get("end"))
        return fin.period(a.get("month") or None)
    except ValueError:
        raise ValueError("日期格式不对（月份 YYYY-MM，日期 YYYY-MM-DD）")


def _wrap(fn):
    try:
        return jsonify(fn())
    except ValueError as e:
        return err(str(e))


@bp.get("/api/fin/meta")
def fin_meta():
    from modules.fin.logic.common import settings
    c = db.conn()
    shops = [dict(r) for r in c.execute("SELECT DISTINCT platform, shop FROM orders WHERE COALESCE(platform,'')!='' ORDER BY platform, shop")]
    return jsonify({"ok": True, "categories": fin.expense_categories(), "settings": settings(), "shops": shops})


def _od_args():
    """订单明细的筛选：期间（month 或 start/end）+ 平台 / 店铺 / 回款状态 / 搜索。"""
    a = request.args
    kw = {"platform": a.get("platform") or "", "shop": a.get("shop") or "", "settle": a.get("settle") or "", "q": (a.get("q") or "").strip()}
    if a.get("start"):
        s, e, _ = fin.period(start=a.get("start"), end=a.get("end"))
        kw.update(start=s, end=e)
    else:
        kw["month"] = a.get("month") or None
    return kw


@bp.get("/api/fin/orders")
def fin_orders():
    """v2.32：订单明细（每单一行：收入、平台费用拆分、到账、成本、利润、回款状态）。"""
    from modules.fin.logic import order_detail
    try:
        return jsonify(order_detail.rows(**_od_args()))
    except ValueError as e:
        return err(str(e))


@bp.get("/api/fin/orders/export")
def fin_orders_export():
    from modules.fin.logic import order_detail
    try:
        buf, label = order_detail.export_xlsx(**_od_args())
    except ValueError as e:
        return err(str(e))
    return send_file(buf, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", as_attachment=True,
                     download_name=f"订单明细_{label.replace(' ~ ', '_')}.xlsx")


@bp.get("/api/fin/overview")
def fin_overview():
    def go():
        if request.args.get("start"):        # v2.41：自定义期间 / 今年
            s, e, _l = _period()
            return fin.overview(start=s, end=e)
        return fin.overview(request.args.get("month") or None)
    return _wrap(go)


@bp.get("/api/fin/pl")
def fin_pl():
    def go():
        s, e, _l = _period()
        return fin.pl(s, e)
    return _wrap(go)


@bp.get("/api/fin/breakdown")
def fin_breakdown():
    def go():
        s, e, label = _period()
        return dict(fin.breakdown(s, e, request.args.get("by") or "platform"), period=label)
    return _wrap(go)


@bp.get("/api/fin/cashflow")
def fin_cashflow():
    def go():
        s, e, label = _period()
        return dict(fin.cashflow(s, e), period=label)
    return _wrap(go)


@bp.get("/api/fin/tax")
def fin_tax():
    def go():
        s, e, label = _period()
        return dict(fin.tax(s, e), period=label)
    return _wrap(go)


@bp.get("/api/fin/reconcile")
def fin_reconcile():
    """v2.49：对账页一次拿全：每个平台的数据状态卡片、数字卡片、每单状态、四类问题、导入记录。platform= 只看一个平台。"""
    def go():
        s, e, label = _period()
        plat = (request.args.get("platform") or "").strip().lower()
        if plat and plat not in ("lazada", "shopee", "tiktok", "other"):
            raise ValueError("平台不对")
        bs = fin.settle_batches(40)
        for b in bs:
            b["source"] = "api" if '"api"' in (b.get("mapping") or "") else "import"
            b.pop("mapping", None)
        return dict(fin.reconcile(s, e, plat), period=label, batches=bs, platforms=fin.platform_status())
    return _wrap(go)


@bp.get("/api/fin/reconcile/order")
def fin_reconcile_order():
    """一单的结算明细按费用分类（佣金、支付手续费、运费、优惠券、罚款…）。"""
    a = request.args
    return _wrap(lambda: fin.order_breakdown((a.get("platform") or "").lower(), a.get("order_no") or ""))


@bp.post("/api/fin/reconcile/check")
def fin_reconcile_check():
    """标「已核对」/ 取消：{platform, order_no, checked: true/false, note}。"""
    b = request.get_json(force=True) or {}
    return _wrap(lambda: fin.set_checked((b.get("platform") or "").lower(), b.get("order_no") or "", bool(b.get("checked", True)),
                                         _uname(), b.get("note") or ""))


@bp.post("/api/fin/settle/sync")
def fin_settle_sync():
    """「立即同步」：Lazada 财务流水（运营分析模块）+ Lazada 结算单（GetPayoutStatus）。"""
    return _wrap(lambda: dict(fin.sync_now(_uname()), ok=True))


# ---------------- 收支 ----------------
@bp.get("/api/fin/expenses")
def fin_expenses():
    def go():
        s, e, label = _period()
        a = request.args
        rows = fin.expenses(s.isoformat(), e.isoformat(), a.get("category") or "", a.get("direction") or "")
        tot_out = sum(r["amount"] for r in rows if r["direction"] == "out")
        tot_in = sum(r["amount"] for r in rows if r["direction"] == "in")
        by = {}
        for r in rows:
            k = (r["direction"], r["category"])
            by[k] = by.get(k, 0) + r["amount"]
        return {"period": label, "rows": rows, "total_out": round(tot_out, 2), "total_in": round(tot_in, 2),
                "by_category": [{"direction": k[0], "category": k[1], "amount": round(v, 2)} for k, v in sorted(by.items(), key=lambda x: -x[1])]}
    return _wrap(go)


@bp.post("/api/fin/expenses")
def fin_expense_save():
    return _wrap(lambda: {"ok": True, "row": fin.save_expense(request.get_json(force=True) or {}, _uname())})


@bp.delete("/api/fin/expenses/<int:eid>")
def fin_expense_delete(eid):
    def go():
        fin.delete_expense(eid, _uname())
        return {"ok": True}
    return _wrap(go)


@bp.post("/api/fin/expenses/<int:eid>/receipt")
def fin_receipt_upload(eid):
    f = request.files.get("file")
    if not f:
        return err("请选择票据照片或 PDF")
    return _wrap(lambda: {"ok": True, "file": fin.set_receipt(eid, f)})


@bp.get("/api/fin/expenses/<int:eid>/receipt")
def fin_receipt(eid):
    p = fin.receipt_path(eid)
    if not p:
        return err("没有票据", 404)
    return send_file(p)


# ---------------- 平台结算导入 ----------------
@bp.post("/api/fin/settle/preview")
def fin_settle_preview():
    f = request.files.get("file")
    if not f:
        return err("请选择平台导出的结算 / 收入报表")
    return _wrap(lambda: fin.settle_preview(f))


@bp.post("/api/fin/settle/import")
def fin_settle_import():
    b = request.get_json(force=True) or {}
    return _wrap(lambda: fin.settle_import(b.get("token"), b.get("mapping") or {}, b.get("platform") or "", b.get("shop") or "", _uname()))


@bp.delete("/api/fin/settle/batches/<int:bid>")
def fin_settle_delete(bid):
    def go():
        fin.delete_batch(bid, _uname())
        return {"ok": True}
    return _wrap(go)


# ---------------- 导出 Excel（给会计） ----------------
@bp.get("/api/fin/export")
def fin_export():
    import openpyxl
    from openpyxl.styles import Font
    try:
        s, e, label = _period()
    except ValueError as ex:
        return err(str(ex))
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "损益表"
    p = fin.pl(s, e)
    ws.append([f"损益表 {p['period']}", "", "说明"])
    for ln in p["lines"]:
        ws.append([("    " if ln["level"] == 1 else "") + ln["zh"], ln["amount"], ("（估）" if ln["est"] else "") + (ln["note"] or "")])
        if ln["level"] >= 2:
            ws.cell(ws.max_row, 1).font = Font(bold=True)
            ws.cell(ws.max_row, 2).font = Font(bold=True)
    for by, title in (("platform", "按平台"), ("shop", "按店铺"), ("product", "按产品"), ("sku", "按SKU"), ("day", "按日")):
        w = wb.create_sheet(title)
        b = fin.breakdown(s, e, by)
        w.append([title, "名称", "订单", "件数", "收入", "平台扣费", "生产成本", "利润", "利润率%"])
        for r in b["rows"]:
            w.append([r["key"], r.get("name") or "", r["orders"], r["qty"], r["revenue"], r["fee"], r["cost"], r["profit"], r["margin"]])
    w = wb.create_sheet("收支明细")
    w.append(["日期", "方向", "分类", "金额", "VAT", "有税票", "收款方", "付款方式", "备注", "记账人"])
    for r in fin.expenses(s.isoformat(), e.isoformat()):
        w.append([r["day"], "收入" if r["direction"] == "in" else "支出", r["category"], r["amount"], r.get("vat") or 0,
                  "是" if r.get("tax_invoice") else "", r.get("payee") or "", _method_zh(r.get("method")), r.get("note") or "", r.get("username") or ""])
    w = wb.create_sheet("现金流")
    cf = fin.cashflow(s, e)
    w.append(["日期", "平台回款", "其他收入", "老板投入", "采购付款", "支出", "净流入", "累计"])
    for r in cf["rows"]:
        w.append([r["key"], r["settle_in"], r["other_in"], r["capital_in"], r["purchase_out"], r["expense_out"], r["net"], r["running"]])
    for sh in wb.worksheets:
        sh.column_dimensions["A"].width = 46 if sh.title == "损益表" else 18
        sh.column_dimensions["B"].width = 16
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"财务_{label}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ---------------- 对外只读接口（密钥权限 finance，规范第 11 条） ----------------
@bp.get("/api/v1/finance/pl")
def v1_fin_pl():
    def go():
        s, e, _l = _period()
        return dict(fin.pl(s, e), ok=True, currency="THB")
    return _wrap(go)


@bp.get("/api/v1/finance/breakdown")
def v1_fin_breakdown():
    def go():
        s, e, label = _period()
        return dict(fin.breakdown(s, e, request.args.get("by") or "platform"), ok=True, period=label, currency="THB")
    return _wrap(go)


@bp.get("/api/v1/finance/cashflow")
def v1_fin_cashflow():
    def go():
        s, e, label = _period()
        return dict(fin.cashflow(s, e), ok=True, period=label, currency="THB")
    return _wrap(go)


@bp.get("/api/v1/finance/expenses")
def v1_fin_expenses():
    def go():
        s, e, label = _period()
        return {"ok": True, "period": label, "currency": "THB", "rows": fin.expenses(s.isoformat(), e.isoformat())}
    return _wrap(go)


@bp.get("/api/v1/finance/tax")
def v1_fin_tax():
    def go():
        s, e, label = _period()
        return dict(fin.tax(s, e), ok=True, period=label, currency="THB")
    return _wrap(go)


# ---- 旧的「其他支出」（利润页上的小表格；v2.30 从 web/orders.py 搬过来）----
from core import finance  # noqa: E402


@bp.get("/api/expenses")
def expenses_list():
    return jsonify({"items": finance.list_expenses()})


@bp.post("/api/expenses")
def expenses_add():
    b = request.get_json(force=True) or {}
    try:
        finance.add_expense(b.get("day"), b.get("kind"), float(b.get("amount")), b.get("note"), _uname())
    except (TypeError, ValueError):
        return err("请填写金额")
    return jsonify({"items": finance.list_expenses()})


@bp.delete("/api/expenses/<int:eid>")
def expenses_delete(eid):
    finance.delete_expense(eid)
    return jsonify({"items": finance.list_expenses()})


# ---------------- v2.53 平台费率 ----------------
@bp.get("/api/fin/fees")
def fin_fees():
    """平台费率页：数字卡片、走势、店铺×类目表、默认费率、Shopee 导入状态、费率变动。"""
    from modules.fin.logic import fees
    return _wrap(lambda: fees.page((request.args.get("platform") or "").strip().lower()))


@bp.post("/api/fin/fees/refresh")
def fin_fees_refresh():
    from modules.fin.logic import fees
    return _wrap(lambda: {"ok": True, "counts": fees.refresh()})


@bp.post("/api/fin/fees/defaults")
def fin_fees_defaults():
    """改默认费率：{platform, rate(0.2=20%)}（平台默认，就是设置里的扣点）或 {platform, category, rate, label}（类目默认；rate 空 = 删除）。"""
    from core import feerates
    b = request.get_json(force=True) or {}
    plat = str(b.get("platform") or "").lower()
    if plat not in feerates.PLATFORMS:
        return err("平台不对")

    def go():
        if b.get("category"):
            feerates.set_category_default(plat, b["category"], b.get("rate"), b.get("label") or "")
        else:
            feerates.set_platform_default(plat, b.get("rate"))
        db.audit(auth.current_user()["username"], "财务·默认费率", plat, str(b.get("category") or "平台默认"))
        return {"ok": True}
    return _wrap(go)


@bp.post("/api/fin/fees/min_orders")
def fin_fees_min_orders():
    from core import feerates
    try:
        n = max(1, int((request.get_json(force=True) or {}).get("min_orders")))
    except (TypeError, ValueError):
        return err("请填整数")
    save_config({"fee_min_orders": n})
    feerates.invalidate()
    return jsonify({"ok": True, "min_orders": n})


@bp.post("/api/fin/fees/recalc")
def fin_fees_recalc():
    """「按新费率重算建议价」：{platform, shop, ids?}。"""
    from core import feerates
    b = request.get_json(force=True) or {}
    plat = str(b.get("platform") or "lazada").lower()
    return _wrap(lambda: dict(feerates.recalc_suggested(load_config(), plat, b.get("shop") or "", b.get("ids"), auth.current_user()["username"]), ok=True))


@bp.post("/api/fin/fees/ack")
def fin_fees_ack():
    from core import feerates
    b = request.get_json(force=True) or {}
    return jsonify({"ok": True, "n": feerates.ack_events(b.get("platform") or None, b.get("shop"), b.get("ids"))})
