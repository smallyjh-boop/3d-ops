# v2.32：财务 → 订单明细（每单的收入、平台扣费拆分、到账、成本、利润、回款状态）和导出
import os as _os
T = _os.environ["T3D_TMP"]; APP = _os.environ["T3D_APP"]
import os, sys, shutil, io, time
D = T + "/m_finorders"; shutil.rmtree(D, ignore_errors=True); shutil.copytree(T + "/v13", D)
import sqlite3; c0 = sqlite3.connect(D + "/app.db"); c0.execute("delete from users"); c0.commit(); c0.close()
os.environ["LH_DATA_DIR"] = D; sys.path.insert(0, APP)
E_ = 0
def ok(b, m):
    global E_; print(("OK  " if b else "ERR ") + m); E_ += (not b)
import app as A
from core import db, config, store
C = A.app.test_client()
config.save_config({"review_enabled": False, "fin_settle_days": 14})
sku = store.load(store.list_products()[0]["id"])["variants"][0]["sku"]
ids = []
for no, tot in (("FO-1", 300), ("FO-2", 200), ("FO-3", 150)):
    ids.append(C.post("/api/orders/manual", json={"order_no": no, "order_total": tot, "items": [{"sku": sku, "qty": 1}]}).get_json()["id"])
c = db.conn(); now = time.time()
c.execute("UPDATE orders SET platform='lazada', shop='Lucky', commission=20, service_fee=10, payment_fee=5, settlement=255 WHERE id=?", (ids[0],))
c.execute("UPDATE orders SET platform='shopee', shop='Lucky', status='shipped', shipped_at=?, order_time=? WHERE id=?", (now - 20 * 86400, now - 20 * 86400, ids[1]))
c.commit()
r = C.get("/api/fin/orders?start=2000-01-01").get_json()
by = {x["order_no"]: x for x in r["rows"]}
ok({"FO-1", "FO-2", "FO-3"} <= set(by), f"rows {len(r['rows'])}")
x = by["FO-1"]
ok(x["settle"] == "settled" and x["fee"] == 45 and x["net"] == 255 and x["fee_source"] == "settlement" and x["commission"] == 20, f"settled row {x['fee']} {x['net']} {x['fee_source']}")
ok(x["profit"] == round(x["revenue"] - x["fee"] - x["cost"], 2), "profit = revenue - fee - cost")
ok(by["FO-2"]["settle"] == "overdue" and by["FO-3"]["settle"] == "pending", "overdue / pending")
ok(r["counts"]["settled"] >= 1 and r["counts"]["overdue"] >= 1, f"counts {r['counts']}")
r2 = C.get("/api/fin/orders?start=2000-01-01&settle=overdue").get_json()
ok(all(x["settle"] == "overdue" for x in r2["rows"]) and r2["rows"], "filter settle")
ok(all(x["platform"] == "lazada" for x in C.get("/api/fin/orders?start=2000-01-01&platform=lazada").get_json()["rows"]), "filter platform")
ok([x["order_no"] for x in C.get("/api/fin/orders?start=2000-01-01&q=FO-3").get_json()["rows"]] == ["FO-3"], "search")
t = r["total"]; ok(abs(t["profit"] - sum(x["profit"] for x in r["rows"])) < 0.05, "totals")
xl = C.get("/api/fin/orders/export?start=2000-01-01")
import openpyxl
wb = openpyxl.load_workbook(io.BytesIO(xl.data)); ws = wb.active
ok(ws["A1"].value == "订单号" and ws.max_row == len(r["rows"]) + 2, f"excel rows {ws.max_row}")
ms = {m["id"]: m for m in C.get("/api/modules?all=1").get_json()["modules"]}
ok(any(mn["tab"] == "fin_orders" for mn in ms["fin"]["menus"]), "menu 订单明细")
C.post("/api/modules/fin", json={"on": False}); ok(C.get("/api/fin/orders").status_code == 404, "404 when fin off"); C.post("/api/modules/fin", json={"on": True})
print("ERRORS", E_); sys.exit(1 if E_ else 0)
