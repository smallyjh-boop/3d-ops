"""导出全部产品资料总表（Excel）：一行一个 SKU + 一行一个产品，方便管理、核对和备份。"""
import datetime as dt

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import codes, pricing

HEAD_FILL = PatternFill("solid", start_color="1F2937")
HEAD_FONT = Font(color="FFFFFF", bold=True)
RED_FILL = PatternFill("solid", start_color="FDE2E2")


def _sheet(wb, title, headers, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for c in ws[1]:
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def build(cfg, products, cat_name_by_id, check_by_id, out_path):
    wb = Workbook()
    wb.remove(wb.active)
    sku_rows, prod_rows = [], []
    from . import stockcalc as _sc
    _av = _sc.available(cfg)
    for p in products:
        codes.ensure_skus(cfg, p)
        prices = pricing.price_product(cfg, p)
        c = p.get("copy", {})
        cat = p.get("category") or {}
        errs, warns = check_by_id.get(p["id"], ([], []))
        status = "可导出" if not errs else f"待补 {len(errs)} 项"
        exported = dt.datetime.fromtimestamp(p["exported_at"]).strftime("%Y-%m-%d") if p.get("exported_at") else ""
        updated = dt.datetime.fromtimestamp(p.get("updated_at") or 0).strftime("%Y-%m-%d %H:%M")
        from . import stockcalc
        stk = stockcalc.stocks(p, cfg, _av)
        for i, v in enumerate(p.get("variants", [])):
            pr = prices[i]
            col = codes.find_color(cfg, v.get("opt1_zh"))
            sku_rows.append([
                v.get("sku"), p["id"], p.get("name_zh", ""), v.get("opt1_zh", ""), col[1] if col else "",
                v.get("opt2_zh", ""), v.get("grams"), v.get("hours"), pr["cost"]["total"], pr["final_price"],
                pr["list_price"], pr["profit"],
                None if pr["actual_markup"] is None else round(pr["actual_markup"] * 100),
                stk[i],
                p.get("material", ""), cat_name_by_id.get(cat.get("id"), ""), c.get("title_th", ""), status,
            ])
        imgs = [i.get("url") for i in p.get("images", []) if i.get("url")]
        pk = p.get("package", {})
        prod_rows.append([
            p["id"], p.get("name_zh", ""), c.get("title_th", ""), c.get("title_en", ""), c.get("title_zh", ""),
            cat_name_by_id.get(cat.get("id"), ""), cat.get("cat_id") or "",
            len(p.get("variants", [])), min((r["final_price"] for r in prices), default=None),
            max((r["final_price"] for r in prices), default=None),
            p.get("material", ""), p.get("size_text", ""),
            f"{pk.get('length')}×{pk.get('width')}×{pk.get('height')}",
            len(imgs), "\n".join(imgs), (p.get("white_bg") or {}).get("url", ""),
            p.get("keywords", ""), p.get("features_zh", ""), status, "；".join(errs), exported, updated,
            _advice_text(p),
        ])
    ws1 = _sheet(wb, "SKU明细", [
        "SKU", "产品编码", "中文名称", "颜色/选项1", "颜色代码", "选项2", "克数", "打印小时", "成本", "售价",
        "划线价", "利润", "加价率%", "库存", "材料", "Lazada类目", "泰文标题", "状态"],
        sku_rows, [22, 11, 22, 12, 9, 10, 8, 9, 9, 9, 9, 9, 9, 7, 8, 26, 50, 12])
    ws2 = _sheet(wb, "产品汇总", [
        "产品编码", "中文名称", "泰文标题", "英文标题", "标题中文回译", "Lazada类目", "类目ID", "变体数",
        "最低价", "最高价", "材料", "尺寸", "包裹cm", "图片数", "图片网址", "白底图", "关键词", "卖点",
        "状态", "待补项目", "首次导出", "更新时间", "打印建议（材质/方向/层高/壁/填充/支撑）"],
        prod_rows, [11, 22, 50, 40, 30, 26, 10, 7, 8, 8, 8, 16, 14, 7, 40, 30, 24, 40, 12, 50, 12, 16, 50])
    for ws, col in ((ws1, 18), (ws2, 19)):
        for row in ws.iter_rows(min_row=2):
            if str(row[col - 1].value or "").startswith("待补"):
                row[col - 1].fill = RED_FILL
    colors = [[zh, code, th, en] for zh, code, th, en in codes.color_table(cfg)]
    _sheet(wb, "颜色代码表", ["中文", "代码", "泰文", "英文"], colors, [12, 8, 14, 18])
    wb.save(out_path)
    return out_path


ORIENT = {"bottom": "底面朝下", "top": "倒放", "front": "前面朝下", "back": "背面朝下", "left": "左侧朝下", "right": "右侧朝下"}
SUPPORT = {"none": "无支撑", "tree": "树状支撑", "normal": "普通支撑"}


def _advice_text(p):
    a = p.get("print_advice") or {}
    if not a:
        return ""
    return (f"{a.get('material', {}).get('pick', '')}；{ORIENT.get(a.get('orientation', {}).get('pick'), '')}；"
            f"层高 {a.get('layer_height_mm')}；壁 {a.get('walls')}；填充 {a.get('infill_pct')}% {a.get('infill_pattern', '')}；"
            f"{SUPPORT.get(a.get('supports'), '')}{'；加帽檐' if a.get('brim') else ''}")
