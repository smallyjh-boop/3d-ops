"""Shopee 批量上传（基础模板，所有类目通用）。

模板：templates_shopee/Shopee_mass_upload_basic_template.xlsx（在 Shopee 卖家中心「批量上传 → 下载模板」获得）
工作表 Template：第 1 行字段代码，第 3 行英文表头，第 7 行起填数据，一行一个 SKU。
类目：core/data/shopee_th_categories.tsv（Shopee 泰国站完整类目树，1442 个叶子类目）。
"""
import datetime as dt
import glob
import html
import io
import os
import re
import zipfile

import openpyxl

from . import codes
from .config import BASE_DIR

TEMPLATE_DIR = os.environ.get("LH_SHOPEE_TEMPLATE_DIR") or os.path.join(BASE_DIR, "templates_shopee")
TREE_PATH = os.path.join(BASE_DIR, "core", "data", "shopee_th_categories.tsv")
DATA_START_ROW = 7
_tree = None

# 常用类目中文名（界面显示用）
ZH_NAMES = {
    "101161": "家居装饰 > 花瓶", "101385": "收藏 > 手办", "101386": "收藏 > 雕像摆件",
    "101181": "园艺 > 花盆", "101258": "收纳 > 桌面收纳", "101254": "收纳 > 收纳盒/篮",
    "101156": "家居装饰 > 相框与墙饰", "101400": "玩具 > 骰子与桌游", "101134": "浴室 > 置物架",
}


def tree():
    """{id: (路径, 是否受限)}"""
    global _tree
    if _tree is None:
        t = {}
        if os.path.exists(TREE_PATH):
            with open(TREE_PATH, encoding="utf-8") as f:
                for line in f:
                    p = line.rstrip("\n").split("\t")
                    if len(p) >= 3:
                        t[p[0]] = (p[1], p[2] == "1")
        _tree = t
    return _tree


# ---------------- v2.25：Shopee 类目属性（卖家中心抓下来的，core/data/shopee_th_attrs.json.gz） ----------------
ATTR_PATH = os.path.join(BASE_DIR, "core", "data", "shopee_th_attrs.json.gz")
_attrs = None


def _attr_data():
    global _attrs
    if _attrs is None:
        import gzip
        import json
        try:
            with gzip.open(ATTR_PATH, "rt", encoding="utf-8") as f:
                _attrs = json.load(f)
        except (OSError, ValueError):
            _attrs = {"attrs": {}, "cats": {}}
    return _attrs


def attrs_for(cid, max_values=120):
    """这个 Shopee 类目的属性：必填的在前。[{id, en, th, required, multi, values: [[id, en, th]], free}]"""
    out = []
    for a in _attr_data()["attrs"].get(str(cid)) or []:
        vals = a.get("v") or []
        out.append({"id": a["id"], "en": a.get("en") or "", "th": a.get("th") or "", "required": bool(a.get("r")),
                    "multi": (a.get("it") in (4, 5)) or (a.get("mx") or 0) > 1, "values": vals[:max_values],
                    "more": max(a.get("more") or 0, len(vals)) > max_values, "free": not vals or a.get("it") in (2, 5)})
    out.sort(key=lambda x: (not x["required"],))
    return out


def required_attrs(cid):
    return [a for a in attrs_for(cid, 0) if a["required"]]


def fill_attrs(cfg, product, limit=8):
    """AI 给这个 Shopee 类目的必填属性（和几个重要属性）挑值。返回 {属性ID: {en, th, value, value_id}}。"""
    from . import ai_copy
    cid = str((product.get("shopee_category") or {}).get("id") or "")
    attrs = attrs_for(cid, 150)
    pick = [a for a in attrs if a["required"]] + [a for a in attrs if not a["required"]][:max(0, limit - sum(1 for a in attrs if a["required"]))]
    if not pick:
        return {}
    import json as _json
    c = product.get("copy") or {}
    info = {"name_zh": product.get("name_zh", ""), "features_zh": (product.get("features_zh") or "")[:1500], "material": product.get("material", ""),
            "size": product.get("size_text", ""), "title_th": c.get("title_th", ""), "category": path_of(cid)}
    spec = [{"id": a["id"], "name": a["en"], "required": a["required"],
             "options": [v[1] for v in a["values"]][:150] if a["values"] else "free text"} for a in pick]
    user = (_json.dumps({"product": info, "attributes": spec}, ensure_ascii=False) +
            '\n\nFor each attribute give the best value for this product. If options are listed you MUST copy one option exactly '
            '(for brand use "No brand" / "NoBrand" style option if present). If unsure and not required, use "". '
            'Never mention 3D printing. Return JSON {"values": {"<attribute id>": "<value>"}}')
    d = ai_copy.extract_json(ai_copy.call(cfg, "You fill Shopee Thailand product attributes. Output only JSON.", user))
    vals = d.get("values") or {}
    out = {}
    for a in pick:
        v = str(vals.get(str(a["id"])) or "").strip()
        if not v:
            continue
        vid = next((x[0] for x in a["values"] if v.lower() in (str(x[1]).lower(), str(x[2]).lower())), None)
        if a["values"] and not a["free"] and vid is None:
            continue                                  # 不在选项里：不要
        th = next((x[2] for x in a["values"] if x[0] == vid), "") if vid else ""
        out[str(a["id"])] = {"en": a["en"], "th": a["th"], "value": v, "value_th": th, "value_id": vid}
    return out


def path_of(cid):
    return (tree().get(str(cid)) or ("", False))[0]


def display_name(cid):
    cid = str(cid)
    return ZH_NAMES.get(cid) or " > ".join(path_of(cid).split(" > ")[-2:])


def search(keywords, limit=60):
    kws = [k.strip().lower() for k in keywords if k and k.strip()]
    scored = []
    for cid, (path, restricted) in tree().items():
        low = path.lower()
        leaf = low.split(" > ")[-1]
        score = sum(3 if k in leaf else 1 for k in kws if k in low)
        if score:
            scored.append((score, cid, path, restricted))
    scored.sort(key=lambda x: (-x[0], x[3], len(x[2])))
    return [(cid, path) for _, cid, path, _r in scored[:limit]]


def template_path():
    os.makedirs(TEMPLATE_DIR, exist_ok=True)
    files = sorted(glob.glob(os.path.join(TEMPLATE_DIR, "*.xlsx")), key=os.path.getmtime)
    return files[-1] if files else None


def _load_workbook(path):
    """Shopee 模板里有不规范的 XML（activePane="bottom_left"），先修正再读。"""
    buf = io.BytesIO()
    with zipfile.ZipFile(path) as zi, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            d = zi.read(it.filename)
            if it.filename.startswith("xl/worksheets/"):
                for a, b in ((b"bottom_left", b"bottomLeft"), (b"bottom_right", b"bottomRight"),
                             (b"top_right", b"topRight"), (b"top_left", b"topLeft")):
                    d = d.replace(b'activePane="' + a + b'"', b'activePane="' + b + b'"')
            zo.writestr(it, d)
    buf.seek(0)
    return openpyxl.load_workbook(buf)


def html_to_text(s):
    """Shopee 详情不支持 HTML，转成纯文本（列表变成 •）。"""
    s = s or ""
    s = re.sub(r"(?i)<li[^>]*>", "• ", s)
    s = re.sub(r"(?i)</li>|<br\s*/?>|</p>|</ul>|</ol>", "\n", s)
    s = re.sub(r"<[^>]+>", "", s)
    s = html.unescape(s)
    s = re.sub(r"[ \t]+\n", "\n", s)
    s = re.sub(r"\n{3,}", "\n\n", s)
    return s.strip()


def _clip(s, n):
    s = (s or "").strip()
    return s if len(s) <= n else s[:n].rstrip()


def check_product(p, prices, cfg):
    """Shopee 专用检查（在通用检查之外）。"""
    errors, warnings = [], []
    c = p.get("copy", {})
    sc = p.get("shopee_category") or {}
    if not sc.get("id"):
        errors.append("还没有选择 Shopee 类目")
    elif sc.get("status") == "check":
        errors.append("AI 对 Shopee 类目不确定，请确认")
    elif tree().get(str(sc["id"]), ("", False))[1]:
        warnings.append("这个 Shopee 类目对你的店铺可能是受限类目，上传可能失败")
    if sc.get("id"):
        req = required_attrs(sc["id"])                # v2.25：必填属性（批量上传的基础模板里没有这几栏）
        have = p.get("shopee_attrs") or {}
        miss = [a["en"] for a in req if str(a["id"]) not in have]
        if miss:
            warnings.append(f"Shopee 这个类目有必填属性：{'、'.join(miss)}（点「AI 填 Shopee 属性」给出建议值；批量上传后在 Shopee 后台补上）")
    title = c.get("title_shopee_th") or c.get("title_th", "")
    if not (20 <= len(title) <= 120):
        if len(title) > 120:
            warnings.append(f"Shopee 标题最多 120 字符，导出时会截断（现在 {len(title)}）")
        else:
            errors.append(f"Shopee 标题需要 20-120 个字符（现在 {len(title)}）")
    desc = html_to_text(c.get("description_th", ""))
    if len(desc) < 60:
        errors.append(f"Shopee 详情至少 60 个字符（现在 {len(desc)}）")
    if not any(i.get("url") for i in p.get("images", [])):
        errors.append("没有图片网址（处理图片后点「上传图床」）")
    vs = p.get("variants", [])
    if len(vs) > 1 or any((v.get("opt1_zh") or "").strip() for v in vs):
        for n, lab in ((1, "变体1名称"), (2, "变体2名称")):
            name = codes.variation_name_th(p, n)
            if name and len(name) > 14:
                errors.append(f"Shopee {lab}最多 14 个字符：「{name}」")
        for v in vs:
            for k in ("opt1_zh", "opt2_zh"):
                if v.get(k):
                    th = codes.option_th(cfg, p, v[k])
                    if len(th) > 20:
                        errors.append(f"Shopee 选项最多 20 个字符：「{th}」")
        imgs = [bool(v.get("image")) for v in vs]
        if any(imgs) and not all(imgs):
            warnings.append("Shopee 要求：给一个颜色配了变体图，其他颜色也都要配，否则变体图会被忽略")
    if not errors:
        from . import lang_guard
        errors += lang_guard.shopee_problems(p, prices, cfg)
    return errors, warnings


def build_rows(products, pricing_by_id, cfg):
    rows = []
    from . import stockcalc as _sc
    _av = _sc.available(cfg)
    for p in products:
        codes.ensure_skus(cfg, p)
        c = p.get("copy", {})
        sc = (p.get("shopee_category") or {}).get("id", "")
        title = _clip(c.get("title_shopee_th") or c.get("title_th", ""), 120)
        desc = html_to_text(c.get("description_th", ""))
        tags = (c.get("hashtags_shopee") or "").strip()
        if tags:  # Shopee 详情里的 #标签 会被搜索
            desc = desc + "\n\n" + tags
        desc = _clip(desc, 5000)
        urls = [i["url"] for i in p.get("images", []) if i.get("url")]
        img_by_raw = {i.get("raw"): i.get("url") for i in p.get("images", [])}
        vs = p.get("variants", [])
        has1 = any((v.get("opt1_zh") or "").strip() for v in vs)
        has2 = any((v.get("opt2_zh") or "").strip() for v in vs)
        all_img = all(img_by_raw.get(v.get("image")) for v in vs) if has1 else False
        pk = p.get("package", {})
        prices = pricing_by_id[p["id"]]
        from . import stockcalc
        stk = stockcalc.stocks(p, cfg, _av)
        for idx, v in enumerate(vs):
            pr = prices[idx]
            vals = {
                "ps_category": int(sc) if str(sc).isdigit() else sc,
                "ps_product_name": title,
                "ps_product_description": desc,
                "ps_sku_parent_short": p["id"],
                "ps_price": pr["final_price"],
                "ps_stock": int(stk[idx]),
                "ps_sku_short": v.get("sku") or f"{p['id']}-{idx + 1:02d}",
                "ps_item_cover_image": urls[0] if urls else "",
                "ps_weight": round(max(0.01, (float(v.get("grams") or 0) + float(cfg.get("packaging_weight_g", 50))) / 1000), 3),
                "ps_length": pk.get("length"),
                "ps_width": pk.get("width"),
                "ps_height": pk.get("height"),
                "channel_id.7000": "On",
            }
            for i, u in enumerate(urls[1:9]):
                vals[f"ps_item_image_{i + 1}"] = u
            if has1:
                vals["et_title_variation_integration_no"] = p["id"]
                vals["et_title_variation_1"] = _clip(codes.variation_name_th(p, 1), 14)
                vals["et_title_option_for_variation_1"] = _clip(codes.option_th(cfg, p, v.get("opt1_zh", "")), 20)
                if all_img:
                    vals["et_title_image_per_variation"] = img_by_raw.get(v.get("image"))
            if has2:
                vals["et_title_variation_2"] = _clip(codes.variation_name_th(p, 2), 14)
                vals["et_title_option_for_variation_2"] = _clip(codes.option_th(cfg, p, v.get("opt2_zh", "")), 20)
            rows.append(vals)
    return rows


def fill(rows, out_path):
    path = template_path()
    if not path:
        raise FileNotFoundError("没有 Shopee 模板（templates_shopee 文件夹）")
    wb = _load_workbook(path)
    ws = wb["Template"]
    col_by_key = {}
    for col in range(1, ws.max_column + 1):
        v = ws.cell(1, col).value
        if v:
            col_by_key[str(v).split("|")[0]] = col
    for r in range(DATA_START_ROW, ws.max_row + 1):
        for cidx in range(1, ws.max_column + 1):
            ws.cell(r, cidx).value = None
    r = DATA_START_ROW
    for vals in rows:
        for k, v in vals.items():
            if v in (None, "") or k not in col_by_key:
                continue
            ws.cell(r, col_by_key[k]).value = v
        r += 1
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb.save(out_path)
    return r - DATA_START_ROW


def export_name():
    return f"shopee_upload_{dt.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
