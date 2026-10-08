"""泰文翻译审核（v2.15）：导出中泰对照表给泰国员工改，改好导入后马上生效（存在 data 里，升级不丢）。

两种翻译：
  - 按编号（static/i18n_keys.js，新功能用，整句翻）
  - 片段替换（static/i18n.js 里的 TH 表，旧功能用：页面上的中文片段换成泰文）
导入的修改存在 data/i18n_overrides.json：{"keys": {编号: 泰文}, "subs": {中文片段: 泰文}}，
首页加载时放进页面（window.I18N_OVERRIDES）。
"""
import io
import json
import os
import re

from .config import BASE_DIR, DATA_DIR

OV_PATH = os.path.join(DATA_DIR, "i18n_overrides.json")
KEYS_JS = os.path.join(BASE_DIR, "static", "i18n_keys.js")
SUBS_JS = os.path.join(BASE_DIR, "static", "i18n.js")
_PAIR = re.compile(r'^\s*"((?:[^"\\]|\\.)*)"\s*:\s*"((?:[^"\\]|\\.)*)",?\s*$')
_KEY = re.compile(r'^\s*"((?:[^"\\]|\\.)*)"\s*:\s*$')
_VAL = re.compile(r'^\s*"((?:[^"\\]|\\.)*)",?\s*$')


def load_overrides():
    try:
        with open(OV_PATH, encoding="utf-8") as f:
            d = json.load(f)
        return {"keys": dict(d.get("keys") or {}), "subs": dict(d.get("subs") or {})}
    except (OSError, ValueError):
        return {"keys": {}, "subs": {}}


def save_overrides(d):
    tmp = OV_PATH + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(d, f, ensure_ascii=False, indent=1)
    os.replace(tmp, OV_PATH)


def script_tag():
    """放进首页的 <script>：导入的翻译修改。"""
    d = load_overrides()
    if not d["keys"] and not d["subs"]:
        return ""
    return "<script>window.I18N_OVERRIDES=" + json.dumps(d, ensure_ascii=False).replace("</", "<\\/") + ";</script>"


def key_entries():
    with open(KEYS_JS, encoding="utf-8") as f:
        txt = f.read()
    body = txt[txt.index("window.I18N_KEYS =") + len("window.I18N_KEYS ="):].strip().rstrip(";")
    return json.loads(body)


def _js_pairs(body):
    """JS 对象字面量里的「"中文": "泰文"」（一行几条、一条两行、注释都认）。"""
    def dec(raw, q):
        if q == "'":
            raw = raw.replace("\\'", "'").replace('"', '\\"')
        return json.loads(f'"{raw}"')
    toks, i, n = [], 0, len(body)
    while i < n:
        ch = body[i]
        if body.startswith("//", i):
            j = body.find("\n", i)
            i = n if j < 0 else j
        elif ch in "\"'":
            j = i + 1
            while j < n and body[j] != ch:
                j += 2 if body[j] == "\\" else 1
            toks.append(("s", dec(body[i + 1:j], ch)))
            i = j + 1
        elif ch == ":":
            toks.append((":", None))
            i += 1
        else:
            i += 1
    out = {}
    for k in range(len(toks) - 2):
        if toks[k][0] == "s" and toks[k + 1][0] == ":" and toks[k + 2][0] == "s":
            out[toks[k][1]] = toks[k + 2][1]
    return out


def sub_entries():
    """片段翻译：i18n.js 里 TH 表的「中文: 泰文」+ 各模块自己的 modules/<id>/i18n/th.json（v2.34 起模块的翻译放在模块里）。"""
    with open(SUBS_JS, encoding="utf-8") as f:
        txt = f.read()
    start = txt.index("const TH = {") + len("const TH = {")
    end = txt.index("\n  };", start)
    out = _js_pairs(txt[start:end])
    try:
        from . import modules
        out.update(modules.i18n_th())
    except Exception:  # noqa: BLE001
        pass
    return out


def export_xlsx():
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    ov = load_overrides()
    wb = openpyxl.Workbook()
    head = Font(bold=True)
    fill = PatternFill("solid", fgColor="FFF3C4")
    ws = wb.active
    ws.title = "按编号"
    ws.append(["编号（不要改）", "中文", "泰文（现在）", "泰文（修改）——只在要改的格子里填"])
    for k, e in key_entries().items():
        ws.append([k, e.get("zh", ""), ov["keys"].get(k) or e.get("th", ""), ""])
    ws2 = wb.create_sheet("片段")
    ws2.append(["中文片段（不要改）", "泰文（现在）", "泰文（修改）——只在要改的格子里填"])
    subs = sub_entries()
    subs.update(ov["subs"])
    for zh, th in subs.items():
        ws2.append([zh, th, ""])
    for w, widths, editcol in ((ws, (22, 50, 50, 50), 4), (ws2, (50, 50, 50), 3)):
        for i, wd in enumerate(widths):
            w.column_dimensions[chr(65 + i)].width = wd
        for c in w[1]:
            c.font = head
        for row in w.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
            row[editcol - 1].fill = fill
        w.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def import_xlsx(data):
    """读「泰文（修改）」一列，非空的记为修改。返回导入条数。"""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True)
    ov = load_overrides()
    keys = key_entries()
    n = 0
    if "按编号" in wb.sheetnames:
        for i, row in enumerate(wb["按编号"].iter_rows(values_only=True)):
            if i == 0 or not row or len(row) < 4:
                continue
            k, new = row[0], row[3]
            if k in keys and new and str(new).strip():
                ov["keys"][k] = str(new).strip()
                n += 1
    if "片段" in wb.sheetnames:
        for i, row in enumerate(wb["片段"].iter_rows(values_only=True)):
            if i == 0 or not row or len(row) < 3:
                continue
            zh, new = row[0], row[2]
            if zh and new and str(new).strip() and re.search(r"[一-鿿]", str(zh)):
                ov["subs"][str(zh)] = str(new).strip()
                n += 1
    if not n and not wb.sheetnames:
        raise ValueError("表格是空的")
    save_overrides(ov)
    return n


def reset():
    save_overrides({"keys": {}, "subs": {}})


def count():
    d = load_overrides()
    return len(d["keys"]) + len(d["subs"])
