"""解析 Lazada 卖家中心导出的批量上新模板，并把产品写进去。

模板结构（每个类目一个工作表）：
  <类目>        第1行中文表头，第2行必填/选填，第3-4行说明，第5行起填数据
  <类目>_hide   第3行字段代码（title.th_TH、sku.price…），第4行字段类型，
                第5行字段属性JSON，第6行下拉选项数量，第7行起是下拉选项
按字段代码而不是列号来填，所以换任何类目的模板都能用。
"""
import datetime as dt
import json
import os
import re

import openpyxl
from openpyxl.utils import get_column_letter

from . import codes

SPLIT = "####splitter line####"
DATA_START_ROW = 5



def _meta(raw):
    try:
        return json.loads(raw) if raw else {}
    except (json.JSONDecodeError, TypeError):
        return {}


def parse(path):
    """返回模板里所有类目工作表及其字段。"""
    wb = openpyxl.load_workbook(path, read_only=False)
    sheets = []
    for ws in wb.worksheets:
        hide_name = ws.title + "_hide"
        if hide_name not in wb.sheetnames:
            continue
        hs = wb[hide_name]
        fields = []
        cat_id = ""
        for col in range(1, hs.max_column + 1):
            key = hs.cell(3, col).value
            if not key:
                continue
            kind = hs.cell(4, col).value or ""
            meta = _meta(hs.cell(5, col).value)
            count = hs.cell(6, col).value
            options, option_ids = [], []
            if count:
                r = 7
                seen_split = False
                while r < 7 + 2 * int(count) + 2:
                    v = hs.cell(r, col).value
                    r += 1
                    if v is None:
                        break
                    if v == SPLIT:
                        seen_split = True
                        continue
                    (option_ids if seen_split else options).append(str(v))
            if len(option_ids) != len(options):
                option_ids = []
            if key == "catId":
                dv = _meta(meta.get("defaultValue"))
                cat_id = str(dv.get("value") or "")
            fields.append({
                "col": col,
                "letter": get_column_letter(col),
                "key": key,
                "label": ws.cell(1, col).value or key,
                "required": ws.cell(2, col).value in ("必填", "บังคับการกรอกข้อมูล", "Mandatory") or bool(meta.get("isRequired")),
                "kind": kind,
                "hidden": bool(meta.get("hide")),
                "readonly": bool(meta.get("readonly")),
                "multi": bool(meta.get("multiSelect")),
                "options": options,
                "option_ids": option_ids,
                "tip": (meta.get("toolTips") or ws.cell(4, col).value or "")[:300],
            })
        existing = 0
        for r in range(DATA_START_ROW, ws.max_row + 1):
            if any(ws.cell(r, c).value not in (None, "") for c in range(1, ws.max_column + 1)):
                existing += 1
        sheets.append({"sheet": ws.title, "cat_id": cat_id, "fields": fields, "existing_rows": existing})
    wb.close()
    return sheets


def label_for_id(field, option_id):
    """按选项 ID 取该模板语言下的选项文字（不同语言的模板 ID 相同）。"""
    ids = field.get("option_ids") or []
    if option_id in ids:
        return field["options"][ids.index(option_id)]
    return None


# 颜色表英文名 → Lazada 颜色系选项 ID
COLOR_FAMILY_IDS = {
    "white": "95204", "black": "4660", "grey": "94920", "gray": "94920", "silver": "4864", "gold": "149729",
    "red": "95308", "pink": "142418", "purple": "95080", "blue": "6831", "sky blue": "149495",
    "green": "94919", "yellow": "95407", "orange": "95049", "brown": "4799", "beige": "3365",
    "clear": "93859", "wood": "4799", "marble": "95204", "glow in the dark": "141068",
    "rainbow": "141068", "multicolor": "141068",
}
MULTICOLOR_ID = "141068"


def attribute_fields(sheet_info):
    """需要你在导出页选一次默认值的类目属性（品牌、材质、保修等）。"""
    out = []
    for f in sheet_info["fields"]:
        k = f["key"]
        if f["hidden"] or f["readonly"]:
            continue
        if k.startswith("productQualifications.") and k.endswith("licenseCode") or \
                (k.startswith("productQualifications.") and f["required"]):
            out.append(f)
            continue
        if k.startswith("catProperty.") or k in ("warrantyType", "warrantyPeriod", "warrantyPolicy",
                                                 "radioDangerousGoods", "deliveryStandard") \
                or f["kind"] == "colorFamily":
            out.append(f)
    return out


def build_rows(products, pricing_by_id, cfg, cat=None):
    """把产品展开成一行一个变体（SKU）。返回 [(product, variant_idx, {key: value})]。"""
    rows = []
    today = dt.datetime.now().replace(microsecond=0)
    from . import stockcalc as _sc
    _av = _sc.available(cfg)
    for p in products:
        codes.ensure_skus(cfg, p)
        ensure_brand(p, cat, cfg)          # 导出时品牌空着也补上 No Brand
        c = p.get("copy", {})
        imgs = [i["url"] for i in p.get("images", []) if i.get("url")][:8]
        white = (p.get("white_bg") or {}).get("url", "")
        variants = p.get("variants", [])
        has1 = any((v.get("opt1_zh") or "").strip() for v in variants)
        has2 = any((v.get("opt2_zh") or "").strip() for v in variants)
        img_by_raw = {i.get("raw"): i.get("url") for i in p.get("images", [])}
        pk = p.get("package", {})
        prices = pricing_by_id[p["id"]]
        attrs = {k: a.get("value") for k, a in (p.get("attributes") or {}).items() if a.get("value") not in (None, "")}
        from . import stockcalc
        stk = stockcalc.stocks(p, cfg, _av)
        for idx, v in enumerate(variants):
            pr = prices[idx]
            vals = dict(attrs)
            vals.update({
                "productNoForBatch": p["id"],
                "title.th_TH": c.get("title_th", ""),
                "title.en_TH": c.get("title_en", ""),
                "marketImages.1:1": white,
                "description": c.get("description_th", ""),
                "enDescription": c.get("description_en", ""),
                "packageContent": c.get("package_content_th", ""),
                "sku.price": pr["list_price"],
                "sku.SellerSku": v.get("sku") or f"{p['id']}-{idx + 1:02d}",
                "sku.package_weight": round(max(0.01, (float(v.get("grams") or 0) + float(cfg.get("packaging_weight_g", 50))) / 1000), 3),
                "sku.package_length": pk.get("length"),
                "sku.package_width": pk.get("width"),
                "sku.package_height": pk.get("height"),
            })
            for i, u in enumerate(imgs):
                vals[f"mainImage.{i}"] = u
            if pr.get("special_price"):
                vals["sku.special_price.SpecialPrice"] = pr["special_price"]
                vals["sku.special_price.Start"] = today.strftime("%Y-%m-%d %H:%M:%S")
                vals["sku.special_price.End"] = (today + dt.timedelta(days=365)).strftime("%Y-%m-%d %H:%M:%S")
            if has1:
                vals["customSaleProp.Custom_Variation_Name_1"] = codes.variation_name_th(p, 1)
                vals["customSaleProp.Custom_Opt_Name_1"] = codes.option_th(cfg, p, v.get("opt1_zh", ""))
                vurl = img_by_raw.get(v.get("image")) if v.get("image") else None
                if vurl:
                    vals["customSaleProp.First_Custom_Opt_Img"] = vurl
                    vals["sku.images.0"] = vurl
            if has2:
                vals["customSaleProp.Custom_Variation_Name_2"] = codes.variation_name_th(p, 2)
                vals["customSaleProp.Custom_Opt_Name_2"] = codes.option_th(cfg, p, v.get("opt2_zh", ""))
            if cat is not None:
                cf = next((f for f in cat["fields"] if f["kind"] == "colorFamily"), None)
                if cf and cf["required"] and not vals.get(cf["key"]):
                    col = codes.find_color(cfg, v.get("opt1_zh"))
                    cid = COLOR_FAMILY_IDS.get((col[3] if col else "").lower(), MULTICOLOR_ID)
                    vals[cf["key"]] = label_for_id(cf, cid) or label_for_id(cf, MULTICOLOR_ID) or ""
            vals["__stock__"] = int(stk[idx])          # 手填 / 默认库存，开了「按耗材自动算库存」时按耗材算（v2.16）
            rows.append((p, idx, vals))
    return rows


def ensure_brand(p, cat, cfg=None):
    """品牌（ยี่ห้อ）：我们所有店铺都是无品牌，空着就自动填设置里的品牌（默认 No Brand）。返回是否改了。"""
    if not cat or not p.get("attributes"):
        return False
    brand = str((cfg or {}).get("brand") or "").strip() or "No Brand"
    changed = False
    for f in cat.get("fields") or []:
        if f.get("kind") != "brandProp" or f.get("hidden") or f.get("readonly"):
            continue
        a = p["attributes"].get(f["key"]) or {}
        if not str(a.get("value") or "").strip():
            p["attributes"][f["key"]] = {"value": brand, "status": "auto", "note": "所有店铺都是无品牌（设置 → 品牌）", "label_zh": "品牌"}
            changed = True
    return changed


def check_product(p, prices, cat=None, cfg=None):
    """单个产品的上架检查。errors 会阻止导出，warnings 只提醒。"""
    errors, warnings = [], []
    c = p.get("copy", {})
    t = c.get("title_th", "")
    if not (5 <= len(t) <= 255):
        errors.append(f"泰文标题需要 5-255 个字符（现在 {len(t)}）")
    if not any(i.get("url") for i in p.get("images", [])):
        errors.append("没有图片网址（处理图片后点「上传图床」）")
    elif any(not i.get("url") for i in p.get("images", [])):
        warnings.append("有图片还没上传图床，导出时会跳过")
    if not p.get("white_bg"):
        warnings.append("没有白底图（选填，有助曝光）")
    d = c.get("description_th", "")
    from .ai_copy import mentions_process
    if any(mentions_process(c.get(k, "")) for k in ("title_th", "title_en", "description_th", "description_en", "package_content_th")):
        warnings.append("文案里提到了 3D 打印/按单生产等生产方式，建议删掉（或重新生成文案）")
    if not d:
        warnings.append("没有泰文详情")
    if c.get("title_th") and not (c.get("title_en") or "").strip():
        warnings.append("没有英文标题（Lazada 有英文标题更好；点「✨ 补全空着的文案」）")
    if not (c.get("package_content_th") or "").strip():
        warnings.append("没有包装内容（What's in the box）")
    if not (c.get("hashtags_shopee") or "").strip():
        warnings.append("没有 Shopee 标签（会被搜索，建议补上）")
    elif not (20 <= len(d) <= 3000):
        warnings.append(f"详情建议 20-3000 字符（现在 {len(d)}）")
    skus = [(v.get("sku") or "").strip().upper() for v in p.get("variants", []) if (v.get("sku") or "").strip()]
    dup = sorted({x for x in skus if skus.count(x) > 1})
    if dup:      # v2.39：重复 SKU 导出后平台订单会对错颜色
        errors.append(f"SKU 重复：{'、'.join(dup)}（每个变体的 SKU 要不一样）")
    pk = p.get("package", {})
    for k, lab in (("length", "长"), ("width", "宽"), ("height", "高")):
        try:
            if not (0.01 <= float(pk.get(k) or 0) <= 200):
                errors.append(f"包裹{lab}需在 0.01-200cm")
        except (TypeError, ValueError):
            errors.append(f"包裹{lab}不是数字")
    if p.get("package_status") == "check":
        errors.append("包裹尺寸是 AI 估算的，请确认（改一下数字或点「确认」）")
    opts = set()
    for i, pr in enumerate(prices):
        v = p["variants"][i]
        if pr["final_price"] <= 0:
            errors.append(f"变体{i + 1}：价格为 0")
        if pr["warning"]:
            warnings.append(f"变体{i + 1}：{pr['warning']}")
        if v.get("grams_src") == "est":
            warnings.append(f"变体{i + 1}：克数/时长是估算值，切片后会更准确")
        combo = (v.get("opt1_zh", ""), v.get("opt2_zh", ""))
        if combo in opts:
            errors.append(f"变体选项重复 {combo}")
        opts.add(combo)
    if len(p.get("variants", [])) > 1 and not any(v.get("opt1_zh") for v in p["variants"]):
        errors.append("有多个变体但没填选项名称")
    if cfg is not None:
        cjk = re.compile(r"[\u4e00-\u9fff]")
        for k in ("opt1_zh", "opt2_zh"):
            for zh in sorted({(v.get(k) or "").strip() for v in p.get("variants", [])} - {""}):
                if cjk.search(codes.option_th(cfg, p, zh)):
                    errors.append(f"选项「{zh}」还没有泰文（在⑤的选项对照里补上，或把颜色加入颜色表）")
        for n, lab in ((1, "变体1名称"), (2, "变体2名称")):
            used = any((v.get(f"opt{n}_zh") or "").strip() for v in p.get("variants", []))
            th = codes.variation_name_th(p, n)
            if used and cjk.search(th):
                errors.append(f"{lab}「{th}」还没有泰文（在⑤里填变体名称的泰文）")
        if p.get("exported_at"):
            for i, v in enumerate(p.get("variants", [])):
                col = codes.find_color(cfg, v.get("opt1_zh"))
                if col and v.get("sku") and col[1] not in v["sku"].split("-"):
                    warnings.append(f"变体{i + 1} 颜色改成了「{v.get('opt1_zh')}」，但 SKU {v['sku']} 已锁定未变")
    catinfo = p.get("category") or {}
    if not catinfo.get("id"):
        errors.append("还没有选择 Lazada 类目")
    elif cat is None:
        errors.append("这个类目的模板还没有（系统会自动下载全部类目模板，进度在「导出上架表格」页；也可以重新选类目）")
    else:
        if catinfo.get("status") == "check":
            errors.append("AI 对类目不确定，请确认类目（下拉框选一下即可）")
        labels = {f["key"]: f["label"] for f in cat["fields"]}
        for k, a in (p.get("attributes") or {}).items():
            if k not in labels:
                continue
            if a.get("status") == "missing":
                errors.append(f"属性「{labels[k]}」必填，请补充")
            elif a.get("status") == "check":
                errors.append(f"属性「{labels[k]}」AI 不确定，请确认")
        if not p.get("attributes"):
            errors.append("类目属性还没填写（点「AI 匹配类目并填写属性」或「按当前类目填写属性」）")
        else:
            attrs = p.get("attributes") or {}
            for f in attribute_fields(cat):
                if f.get("required") and f["key"] not in attrs:
                    errors.append(f"属性「{f.get('label', f['key'])}」必填，还没填写（点「按当前类目填写属性」）")
    if cfg is not None and not errors:
        # 最后把要导出的每个字段查一遍：平台上只能是泰文 / 英文（v2.19.1）
        from . import lang_guard
        errors += lang_guard.lazada_problems(p, prices, cfg, cat)
    ir = p.get("ip_risk") or {}
    if ir.get("status") == "pending" and ir.get("items"):
        errors.append("侵权待确认：选品工具提示可能涉及品牌 / 角色（" + "、".join(x.get("term", "") for x in ir["items"][:5])
                      + "），在「⑤ 上架文案」里确认后才能导出")
    from . import specs as _specs
    e2, w2 = _specs.review_problems(p)
    errors += e2
    warnings += w2
    return errors, warnings


def fill_multi(template_path, rows_by_sheet, out_path):
    """在同一个模板文件里，按工作表（类目）分别写入产品行。"""
    wb = openpyxl.load_workbook(template_path)
    infos = {s["sheet"]: s for s in parse(template_path)}
    total = 0
    for sheet_name, rows in rows_by_sheet.items():
        ws = wb[sheet_name]
        info = infos[sheet_name]
        col_by_key = {f["key"]: f["col"] for f in info["fields"]}
        valid = {f["key"]: f for f in info["fields"]}
        stock_cols = [f["col"] for f in info["fields"] if f["key"].startswith("sku.multiWarehouseQty.")]
        for r in range(DATA_START_ROW, ws.max_row + 1):
            for cidx in range(1, ws.max_column + 1):
                ws.cell(r, cidx).value = None
        r = DATA_START_ROW
        for _p, _idx, vals in rows:
            for k, v in vals.items():
                if k.startswith("__") or v in (None, ""):
                    continue
                f = valid.get(k)
                if not f or f["readonly"]:
                    continue
                ws.cell(r, col_by_key[k]).value = v
            for col in stock_cols:
                ws.cell(r, col).value = vals["__stock__"]
            r += 1
        total += r - DATA_START_ROW
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb.save(out_path)
    return total
