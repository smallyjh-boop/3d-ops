# v2.15：用真实数据校准（克数系数、打印速度、切片 vs 实际时长、空打阈值）+ 按编号翻译导出导入
import os as _os
T = _os.environ["T3D_TMP"]; APP = _os.environ["T3D_APP"]
import os, sys, shutil, io, time, json
D=T+"/vcal"; shutil.rmtree(D, ignore_errors=True); shutil.copytree(T+"/v13", D)
import sqlite3; c0=sqlite3.connect(D+"/app.db"); c0.execute("delete from users"); c0.commit(); c0.close()
os.environ["LH_DATA_DIR"]=D; sys.path.insert(0,APP)
import app as A
from core import store, db, slicer, calib, config, print_files, production, modellib
c=A.app.test_client(); E=0
def ok(b,m):
    global E; print(("OK  " if b else "ERR ")+m); E+=(not b)
print_files.init(); production.init() if hasattr(production,"init") else None
cfg=config.load_config()
pids=[p["id"] for p in store.list_products()][:6]
mets=[]
for pid in pids:
    p=store.load(pid); p["material"]="PLA"
    p["model"]={"file":"x.stl","metrics":{"volume_mm3":20000+5000*len(mets),"area_mm2":6000,"size":[40,40,40]}}
    store.save(p); mets.append(p["model"]["metrics"])
cn=db.conn()
P1=(cfg["printers"][0].get("name") or "")
for i,pid in enumerate(pids[:5]):
    est=slicer.estimate(dict(cfg,est_calibration=1.0,est_calibration_by_mat={}),mets[i],"PLA")["grams"]
    cn.execute("INSERT INTO print_files(product_id,printer,grams,hours,per_plate,active,status,kind,created_at) VALUES(?,?,?,?,?,?,?,?,?)",
               (pid,P1,est*1.2*2,(est*1.2*2)/30+0.15,2,1,"reviewed","fdm",time.time()))
cn.commit()
r=c.get("/api/calib").get_json()
g=[x for x in r["grams"] if x["material"]=="PLA"][0]
ok(g["n"]==5 and abs(g["suggest"]-1.2)<0.02, f"grams factor suggest {g}")
sp=r["speeds"][0]; ok(sp["n"]>=5 and abs(sp["suggest"]-30)<0.5, f"speed suggest {sp}")
# apply
before=slicer.estimate(config.load_config(),mets[0],"PLA")["grams"]
r2=c.post("/api/calib/apply",json={"kind":"grams","data":{"PLA":g["suggest"]}}); ok(r2.status_code==200, "apply grams")
after=slicer.estimate(config.load_config(),mets[0],"PLA")["grams"]
ok(abs(after/before-1.2/ (config.load_config().get("est_calibration") or 1))<0.05 or abs(after/before-1.2)<0.05, f"estimate uses factor {before}->{after}")
ok(slicer.estimate(config.load_config(),mets[0],"PETG")["grams"]>0, "other material unaffected path ok")
ok(c.post("/api/calib/apply",json={"kind":"grams","data":{"PLA":9}}).status_code==400, "absurd factor refused")
c.post("/api/calib/apply",json={"kind":"speed","data":{"0":sp["suggest"]}})
ok(abs(config.load_config()["printers"][0]["g_per_h"]-sp["suggest"])<0.01, "speed applied to printer 0")
# actual time vs sliced
now=time.time()
for i in range(4):
    cn.execute("INSERT INTO print_runs(printer,product_id,sku,count,grams,hours,status,started_at,ended_at) VALUES(?,?,?,?,?,?,?,?,?)",
               (0,pids[0],"X",1,10,2.0,"done",now-9000,now-9000+2.0*1.1*3600))
cn.commit()
t0=[x for x in c.get("/api/calib").get_json()["times"] if x["idx"]==0][0]; ok(t0["n"]==4 and abs(t0["ratio"]-1.1)<0.01, f"time ratio {t0}")
# failures vs orientation: product with model_lib orient level ok but adhesion failures
p=store.load(pids[1]); p["model_lib"]=[{"id":"m1","info":{"orients":[{"key":"bottom","level":"ok","contact_cm2":2.2}]}}]; p["model"]["lib_id"]="m1"; p["recipe"]={"orientation":"bottom"}; store.save(p)
for i in range(12):
    cur=cn.execute("INSERT INTO print_runs(printer,product_id,sku,count,grams,hours,status,started_at,ended_at) VALUES(?,?,?,?,?,?,?,?,?)",
                   (0,pids[1],"Y",1,10,1,"failed" if i<4 else "done",now-3600,now-1800))
    if i<4: cn.execute("INSERT INTO failures(ts,run_id,printer,product_id,sku,reason,grams) VALUES(?,?,?,?,?,?,?)",(now,cur.lastrowid,0,pids[1],"Y","打飞 / 脱离热床（空打）",5))
cn.commit()
r=c.get("/api/calib").get_json()
okst=r["failures"]["ok"]; ok(okst["runs"]>=12 and okst["adhesion"]==4, f"failure stats ok-level {okst}")
ok(r["risk_suggest"] and r["risk_suggest"]["contact_warn"]>=2.2, f"risk suggestion {r['risk_suggest']}")
c.post("/api/calib/apply",json={"kind":"risk","data":{"contact_warn":2.5}})
ok(calib.risk_thresholds()["contact_warn"]==2.5, "risk threshold saved")
st={"contact_cm2":2.0,"overhang_cm2":0,"height_mm":10,"tall_ratio":1}
ok(any(x["code"]=="risk.contact_warn" for x in modellib.orient_risks(st)), "new threshold used by risk rules")
c.post("/api/calib/apply",json={"kind":"risk","data":"reset"})
ok(calib.risk_thresholds()["contact_warn"]==1.5 and not modellib.orient_risks(st), "risk reset")
ok("production" not in (c.get("/api/calib").get_json() or {}) , "report json")
# i18n export / import
x=c.get("/api/system/i18n/export"); ok(x.status_code==200 and len(x.data)>10000, "i18n export")
import openpyxl
wb=openpyxl.load_workbook(io.BytesIO(x.data)); ws=wb["按编号"]; ws2=wb["片段"]
row=[r_ for r_ in ws.iter_rows(min_row=2) if r_[0].value=="pv.fits"][0]; row[3].value="วางได้นะ"
ws2.cell(row=2,column=3).value="แก้แล้ว"; zh2=ws2.cell(row=2,column=1).value
b=io.BytesIO(); wb.save(b)
r=c.post("/api/system/i18n/import",data={"file":(io.BytesIO(b.getvalue()),"t.xlsx")},content_type="multipart/form-data").get_json()
ok(r.get("imported")==2, f"i18n import {r}")
html=c.get("/").data.decode()
ok("I18N_OVERRIDES" in html and "วางได้นะ" in html and json.dumps(zh2,ensure_ascii=False)[1:-1] in html, "overrides injected into page")
c.post("/api/system/i18n/reset",json={}); ok("I18N_OVERRIDES" not in c.get("/").data.decode(), "i18n reset")
print("ERRORS",E)
